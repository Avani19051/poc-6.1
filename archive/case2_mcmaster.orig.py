"""
Case 2 - McMaster-Carr Material / RoHS lookup
Language: Python 3.9+   (browser automation: Playwright + Chromium)

1. Asks the user for an MPN (must be one of the MPNs in TableA.xlsx).
2. Opens https://www.mcmaster.com/, searches the MPN, clicks "Product Detail"
   when the results page shows it, and waits for the specification table.
3. Reads "Material" and "RoHS Compliance" from the product detail page.
4. Writes both values into that MPN's row of TableA and saves the complete
   table as  TableA_Updated.xlsx  in the user's Documents folder.

Usage:
    python case2_mcmaster.py              # interactive: prompts for MPNs
    python case2_mcmaster.py 91864A028    # one or more MPNs from the command line
    python case2_mcmaster.py --all        # every MPN in TableA
    python case2_mcmaster.py --headed     # show the browser window (useful if blocked)

Install once:
    pip install playwright openpyxl
    playwright install chromium
"""

from __future__ import annotations

import argparse
import logging
import random
import re
import sys
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

log = logging.getLogger("case2")

BASE_URL = "https://www.mcmaster.com/"
HERE = Path(__file__).parent
TABLE_A = HERE / "TableA.xlsx"
OUTPUT_NAME = "TableA_Updated.xlsx"
HEADERS = ["MPN", "MANUFACTURER NAME", "Material", "RoHS"]
DEFAULT_ROWS = [
    ["91864A028", "MCMASTER CARR SUPPLY", "?", "?"],
    ["90044A249", "MCMASTER CARR SUPPLY", "?", "?"],
    ["91251A107", "MCMASTER CARR SUPPLY", "?", "?"],
    ["91251A105", "MCMASTER CARR SUPPLY", "?", "?"],
]
MPN_PATTERN = re.compile(r"^\d{4,5}[A-Z]\d{1,4}$")   # McMaster style, e.g. 91864A028

PAGE_TIMEOUT_MS = 45_000
RETRIES = 3


# --------------------------------------------------------------------------- #
# Table A (Excel) helpers
# --------------------------------------------------------------------------- #
def documents_dir() -> Path:
    home = Path.home()
    for candidate in (home / "OneDrive" / "Documents", home / "Documents"):
        if candidate.is_dir():
            return candidate
    (home / "Documents").mkdir(parents=True, exist_ok=True)
    return home / "Documents"


def ensure_table_a(path: Path = TABLE_A) -> None:
    """Create TableA.xlsx with the rows from the brief if it does not exist."""
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "TableA"
    ws.append(HEADERS)
    for row in DEFAULT_ROWS:
        ws.append(row)
    style_sheet(ws)
    wb.save(path)
    log.info("Created %s", path)


def style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(
            14, max(len(str(c.value or "")) for c in col) + 2)


def load_table(input_path: Path, output_path: Path):
    """Continue from a previous run's output if present, so earlier results are kept."""
    source = output_path if output_path.exists() else input_path
    wb = load_workbook(source)
    ws = wb.active
    header = [str(c.value).strip().upper() for c in ws[1]]
    try:
        cols = {h: header.index(h.upper()) + 1 for h in HEADERS}
    except ValueError as exc:
        raise SystemExit(f"TableA is missing an expected column: {exc}")
    rows = {}
    for r in range(2, ws.max_row + 1):
        mpn = ws.cell(r, cols["MPN"]).value
        if mpn:
            rows[str(mpn).strip().upper()] = r
    return wb, ws, cols, rows


# --------------------------------------------------------------------------- #
# Scraping
# --------------------------------------------------------------------------- #
def parse_specs(page_text: str) -> dict[str, str]:
    """
    McMaster's spec table renders as alternating label / value lines:
        Material
        Black-Oxide Alloy Steel
        RoHS Compliance
        RoHS 3 (2015/863/EU) Compliant
    """
    lines = [ln.strip() for ln in page_text.splitlines() if ln.strip()]
    specs: dict[str, str] = {}
    for label in ("Material", "RoHS Compliance"):
        for i, ln in enumerate(lines[:-1]):
            if ln.lower() == label.lower():
                specs[label] = lines[i + 1]
                break
    return specs


class McMasterScraper:
    def __init__(self, headed: bool = False):
        from playwright.sync_api import sync_playwright  # imported lazily for nicer errors
        self._pw = sync_playwright().start()
        self.browser = self._pw.chromium.launch(headless=not headed)
        self.context = self.browser.new_context(
            viewport={"width": 1366, "height": 900},
            user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
        )
        self.page = self.context.new_page()
        self.page.set_default_timeout(PAGE_TIMEOUT_MS)

    def close(self) -> None:
        self.context.close()
        self.browser.close()
        self._pw.stop()

    def _search(self, mpn: str) -> None:
        page = self.page
        page.goto(BASE_URL, wait_until="domcontentloaded")
        box = page.get_by_placeholder(re.compile("search", re.I)).first
        box.wait_for(state="visible")
        box.fill(mpn)
        box.press("Enter")

    def _open_product_detail(self, mpn: str) -> None:
        """Click 'Product Detail' (the red box in the brief) when the page offers it."""
        page = self.page
        detail = page.get_by_text(re.compile(r"^\s*Product Detail\s*$", re.I)).first
        try:
            detail.wait_for(state="visible", timeout=10_000)
            detail.click()
            log.info("%s: clicked 'Product Detail'", mpn)
        except Exception:
            # McMaster often jumps straight to the detail view for an exact part number
            log.info("%s: no 'Product Detail' button - page already shows the detail view", mpn)

    def lookup(self, mpn: str) -> dict[str, str]:
        last_error: Exception | None = None
        for attempt in range(1, RETRIES + 1):
            try:
                self._search(mpn)
                self._open_product_detail(mpn)
                # wait until the specification table has rendered
                self.page.get_by_text("RoHS Compliance", exact=True).first.wait_for(state="visible")
                text = self.page.locator("main").first.inner_text()
                if mpn.upper() not in text.upper():
                    raise RuntimeError("page loaded but does not show this part number")
                specs = parse_specs(text)
                if "Material" not in specs and "RoHS Compliance" not in specs:
                    raise RuntimeError("specification table not found")
                return specs
            except Exception as exc:
                last_error = exc
                wait = 5 * attempt + random.uniform(0, 3)   # back off: site rate-limits fast requests
                log.warning("%s: attempt %d/%d failed (%s); retrying in %.0fs",
                            mpn, attempt, RETRIES, str(exc).splitlines()[0], wait)
                time.sleep(wait)
        raise RuntimeError(f"could not read {mpn} after {RETRIES} attempts: {last_error}")


# --------------------------------------------------------------------------- #
# Main flow
# --------------------------------------------------------------------------- #
def ask_for_mpns(valid: list[str]) -> list[str]:
    print("\nMPNs in TableA:", ", ".join(valid))
    while True:
        raw = input("Enter an MPN (or 'all', or 'q' to quit): ").strip().upper()
        if raw in {"Q", "QUIT", "EXIT", ""}:
            return []
        if raw == "ALL":
            return valid
        if not MPN_PATTERN.match(raw):
            print(f"  '{raw}' does not look like a McMaster part number (e.g. 91864A028).")
        elif raw not in valid:
            print(f"  '{raw}' is not in TableA. Please choose one of: {', '.join(valid)}")
        else:
            return [raw]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mpns", nargs="*", help="MPN(s) to look up (default: ask)")
    ap.add_argument("--all", action="store_true", help="look up every MPN in TableA")
    ap.add_argument("--headed", action="store_true", help="show the browser window")
    ap.add_argument("--table", type=Path, default=TABLE_A, help="input TableA.xlsx")
    ap.add_argument("--out", type=Path, default=None, help="output folder (default: ~/Documents)")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    ensure_table_a(args.table)
    out_path = (args.out or documents_dir()) / OUTPUT_NAME
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb, ws, cols, rows = load_table(args.table, out_path)
    valid = list(rows)

    if args.all:
        targets = valid
    elif args.mpns:
        targets = [m.strip().upper() for m in args.mpns]
        bad = [m for m in targets if m not in rows]
        if bad:
            print(f"Not in TableA: {', '.join(bad)}. Valid MPNs: {', '.join(valid)}")
            targets = [m for m in targets if m in rows]
    else:
        targets = ask_for_mpns(valid)
    if not targets:
        print("Nothing to do.")
        return 0

    try:
        scraper = McMasterScraper(headed=args.headed)
    except Exception as exc:
        print(f"Could not start the browser ({exc}).\nRun:  pip install playwright  &&  playwright install chromium")
        return 1

    failures = 0
    try:
        for i, mpn in enumerate(targets):
            if i:
                time.sleep(random.uniform(4, 8))   # polite gap between lookups
            try:
                specs = scraper.lookup(mpn)
                material = specs.get("Material", "Not Found")
                rohs = specs.get("RoHS Compliance", "Not Found")
                print(f"[OK] {mpn}: Material = {material} | RoHS = {rohs}")
            except Exception as exc:
                failures += 1
                material = rohs = "Lookup failed"
                log.error("%s: %s", mpn, exc)
            r = rows[mpn]
            ws.cell(r, cols["Material"], material)
            ws.cell(r, cols["RoHS"], rohs)
            style_sheet(ws)
            wb.save(out_path)          # save after every MPN so nothing is lost on a crash
    finally:
        scraper.close()

    print(f"\nSaved: {out_path}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
